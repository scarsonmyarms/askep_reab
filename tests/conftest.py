"""
Спільні фікстури.

* `mini_dir` — невеликий синтетичний довідник, створюється в тимчасовій папці. Тести на ньому
  не залежать від реального файлу і перевіряють логіку на контрольованих даних.
* `real_dir` — реальний довідник (Додаток). Шлях задається змінною середовища REHAB_XLSX
  або опцією `pytest --xlsx=шлях`; за замовчуванням шукається dod1.xlsx (у корені проєкту або в tests\\). Якщо файл не знайдено, тести з позначкою `real` пропускаються.
"""
import os
from pathlib import Path

import pytest
from openpyxl import Workbook

from coding_rules import Directory

FO = '"Діагноз_ФО"'
PP = '"Діагноз_ПП"'
SP = '"Діагноз_СП"'
C_ = '"Діагноз_С"'
POST_ONLY = "Так, протягом післягострого реабілітаційного періоду"
LONG_ONLY = "Так, в довготривалому періоді"

# (назва в B, категорії, стовпець N, формула O)
MINI_ROWS = [
    ("I63.3 Інфаркт головного мозку", "ПП", POST_ONLY, f"плюс код(и) {FO} та/або код(и) {C_}"),
    ("I69.3 Наслідки інфаркту мозку", "ПП", LONG_ONLY, f"плюс код(и) {FO}"),
    ("G81.9 Геміплегія, неуточнена", "ФО", "так", f"плюс код {PP}"),
    ("G81.9 Геміплегія, неуточнена", "ФО", "так", f"плюс код {PP}"),               # дублікат
    ("G93.7 Синдром Рейє (Рея)", "ПП,ФО", "так", f"плюс код(и) {FO}"),
    ("G93.1 Аноксичні ураження головного мозку", "ФО", "так",
     f"плюс код {PP} плюс код(и) {FO} та/або код(и) {SP}"),
    ("G11 Спадкова атаксія", "ПП,ФО", "так",
     f"плюс код {PP} та/або код(и) {FO} та/або код(и) {SP}"),
    ("A32.1 Лістеріозний менінгіт та менінгоенцефаліт", "ПП", "так", f"плюс код(и) {FO}"),
    ("A06.6\xa0 Амебний абсцес головного мозку", "ПП", "так", f"плюс код(и) {FO}"),
    ("I46.0 Зупинка серця з вдалим відновленням серцевої діяльності", "ПП", "ні", f"плюс код(и) {FO}"),
    ("R26.2 Утруднення при ходьбі", "СП", "ні", f"плюс код {PP}"),
    ("R47 Розлади мови (R47.0- R47.8)", "ФО,СП", "ні", f"плюс код {PP}"),
    ("Z96.6 Наявність ортопедичних імплантатів суглобів", "С", "ні", f"плюс код {PP}"),
    ("Z50.0 Реабілітація при хворобах серця", "СФЗ", "ні", ""),
    ("Z74.1 Потреба у допомозі при самообслуговуванні", "СФЗ", "ні", f"плюс код {PP}"),
    ("M54.1 Радикулопатія", "ПП", "так, якщо етіологію не встановлено", f"плюс код(и) {FO}"),
    ("M51.1 Порушення поперекових міжхребцевих дисків з радикулопатією (G55.1)", "ПП", "так",
     f"плюс код(и) {FO}"),
    ("G55.1* Компресія нервових корінців при ураженнях міжхребцевого диска", "ФО", "ні", f"плюс код {PP}"),
    ("A17.8† Інший туберкульоз нервової системи G63.0*", "ПП", "так", f"плюс код(и) {FO}"),
    ("G63.0* Поліневропатія при інфекційних хворобах (A17.8 †)", "ФО", "ні", f"плюс код {PP}"),
    ("G01* Менінгіт при бактеріальних хворобах", "ПП", "так", f"плюс код(и) {FO}"),
    ("G55.0* Компресія корінців при новоутвореннях (C00-D48 †)", "ФО", "ні", f"плюс код {PP}"),
    ("C72.0 Злоякісне новоутворення спинного мозку", "ПП", "так", f"плюс код(и) {FO}"),
    ("T90.5 Наслідки внутрішньочерепної травми", "ПП", LONG_ONLY, f"плюс код S06 плюс код(и) {FO}"),
    ("S06.0 Струс головного мозку", "ПП", POST_ONLY, f"плюс код(и) {FO}"),
    ("M13.8 Інший уточнений артрит", "ПП", "так", f"плюс код(и) {FO}"),
    ("M05.0 Ревматоїдний артрит", "ПП", "так", f"плюс код(и) {FO}"),
    ("M45.0 Анкілозуючий спондиліт", "ПП", "так",
     f"плюс код {PP} та/або код(и) {FO} та/або код(и) {SP} або код(и) {C_}"),
]


def make_xlsx(path: Path, rows=MINI_ROWS) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Аркуш1"
    ws["A1"] = "Довідник (тестовий)"
    hdr = ["Код", "Діагноз", "", "", "", "", "", "",
           "Діагноз_ПП", "Діагноз_ФО", "Діагноз_С", "Діагноз_СП", "Діагноз_СФЗ",
           "Основний діагноз", "Формула для кодування"]
    for c, h in enumerate(hdr, 1):
        ws.cell(4, c, h or None)
    for i, (name, cats, main, formula) in enumerate(rows, 5):
        ws.cell(i, 1, name.split()[0])
        ws.cell(i, 2, name)
        cs = cats.split(",") if cats else []
        for col, cat in zip(range(9, 14), ["ПП", "ФО", "С", "СП", "СФЗ"]):
            if cat in cs:
                ws.cell(i, col, "так")
        ws.cell(i, 14, main)
        ws.cell(i, 15, formula)
    wb.save(path)
    return path


@pytest.fixture(scope="session")
def mini_dir(tmp_path_factory) -> Directory:
    return Directory(str(make_xlsx(tmp_path_factory.mktemp("dir") / "mini.xlsx")))


def pytest_addoption(parser):
    parser.addoption("--xlsx", action="store", default=None, help="шлях до реального довідника")
    parser.addoption("--show-reports", action="store_true", default=False,
                     help="у підсумку показати повний звіт перевірки для кожної комбінації")


# ---------------------------------------------------------------------------------------------
# Журнал комбінацій діагнозів, перевірених на реальному довіднику (виводиться в кінці запуску)
# ---------------------------------------------------------------------------------------------
_COMBOS: list[dict] = []


@pytest.fixture
def combo_log():
    """Тест додає сюди запис про кожну перевірену комбінацію (до assert — щоб потрапили й невдалі)."""
    return _COMBOS


def pytest_terminal_summary(terminalreporter, exitstatus, config):
    if not _COMBOS:
        return
    tr = terminalreporter
    tr.section("Комбінації діагнозів, перевірені на реальному довіднику")
    for n, c in enumerate(_COMBOS, 1):
        mark = "OK " if c["passed"] else "НЕ ЗБІГЛОСЯ"
        tr.write_line(f"{n:>2}. [{mark}] {c['title']}")
        for line in c["lines"]:
            tr.write_line(f"      {line}")
        if config.getoption("--show-reports") and c.get("report"):
            for line in c["report"].splitlines():
                tr.write_line(f"        | {line}")
        tr.write_line("")


@pytest.fixture(scope="session")
def real_dir(request) -> Directory:
    here = Path(__file__).resolve().parent
    folders = [here, here / "tests"]          # працює і з кореня проєкту, і з tests\
    candidates = [request.config.getoption("--xlsx"), os.environ.get("REHAB_XLSX")]
    candidates += [str(f / "dod1.xlsx") for f in folders]
    for f in folders:
        candidates += [str(p) for p in f.glob("*.xlsx") if not p.name.startswith("~$")]
    for p in candidates:
        if p and Path(p).is_file():
            return Directory(p, nszu_files=None)      # лише довідник (без файлів НСЗУ)
    pytest.skip(f"Реальний довідник не знайдено. Перевірено: {[c for c in candidates if c]}")


@pytest.fixture(scope="session")
def real_nszu(real_dir) -> Directory:
    """Реальний довідник + файли НСЗУ з тієї самої папки (перелік основних, формули, пари †/*)."""
    d = Directory(real_dir.path)
    if not d.nszu or not all([d.nszu.main_list, d.nszu.formulas_main, d.nszu.formulas_add, d.nszu.pairs]):
        pytest.skip("Файли НСЗУ не знайдено поруч із довідником (потрібні всі 4)")
    return d


@pytest.fixture
def d(mini_dir):
    """Коротке ім'я для пошуку в синтетичному довіднику."""
    return mini_dir
