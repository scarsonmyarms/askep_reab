"""
Логіка кодування діагнозів реабілітаційного випадку.

Джерела:
  * Додаток до наказу НСЗУ (таблиця xlsx): стовпці B (діагноз), I–M (ПП, ФО, С, СП, СФЗ),
    N (чи може бути основним), O (формула кодування);
  * Правила щодо визначення послуг за напрямом реабілітаційної допомоги
    (наказ НСЗУ № 182 від 16.04.2026) — період, подвійне кодування †/*, несумісні поєднання.

Правила застосування формул:
  * формула (O) застосовується лише до основного діагнозу та супутнього 1;
  * список супутнього 1 — за формулою основного; списки супутніх 2..N — за формулами основного
    та супутнього 1 (плюс парні коди †/* уже обраних супутніх);
  * у список потрапляють лише категорії ще не закритих вимог (групи, закриті категорією самого
    діагнозу або вже обраними діагнозами, не враховуються); коли все закрито — показуються всі
    категорії формул, які діагнози не закривають самі;
  * умови формули супутнього 1 закриваються лише ним самим (своєю категорією) та іншими супутніми —
    основний діагноз не враховується;
  * кожне «плюс» — окрема обов'язкова група; «та/або» і «або» — достатньо одного з варіантів
    (так формули записано в нових файлах НСЗУ: «Діагноз_ПП AND (Діагноз_ФО OR Діагноз_С)»);
  * у переліках кодів «та» — окрема обов'язкова група (напр., T92.3: S43/S53/S63 і T11.2);
  * у кожному випадку обов'язковий код «Діагноз_Z» (Z50.x); СФЗ — окрема необов'язкова категорія;
  * категорії кодів — з файлу НСЗУ «Перелік діагнозів ПП, ФО, С, СП, Z» (якщо є поруч із довідником);
  * група вважається виконаною самим діагнозом, якщо він має відповідну категорію (I–M)
    (напр., G93.7 з категоріями ПП, ФО і формулою «плюс ФО» додаткових діагнозів не потребує).

Запуск з командного рядка (без графічного інтерфейсу):
  python coding_rules.py довідник.xlsx --period post --main I63.3 --comp G81.9 --comp R47
"""
from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from typing import Iterable

from functools import cached_property
from pathlib import Path

from openpyxl import load_workbook

import nszu_sources as nszu

CATS = ("ПП", "ФО", "С", "СП", "СФЗ", "Z")   # Z — Діагноз_Z (коди Z50.x), окремо від СФЗ
SETTING_INPATIENT = "Стаціонар"
SETTING_OUTPATIENT = "Амбулаторія"
AMBULATORY_FILE = Path(__file__).with_name("ambulatory_services.json")
PERIOD_POST = "Післягострий"
PERIOD_LONG = "Довготривалий"

OK, WARN, FAIL, INFO = "✔", "⚠", "✖", "ℹ"

# ----------------------------------------------------------------------------------------
# Несумісні поєднання (Правила CR_3_5, CR_4_3, CR_6_3 / табл. 1). Можна доповнювати.
# (префікс коду 1, префікс коду 2, серйозність, правило, пояснення)
# ----------------------------------------------------------------------------------------
EXCLUSIONS: list[tuple[str, str, str, str, str]] = []
for _b in ["M51.1", "M47.2", "M50.1", "S24.2", "S14.2", "S34.2"]:
    EXCLUSIONS.append(("M54.1", _b, FAIL, "CR_3_5",
                       "Не слід одночасно застосовувати M54.1 з кодом, що вже містить ознаку радикулопатії / "
                       "ураження корінця; перевага — коду, що відображає етіологію та обмеження."))
EXCLUSIONS.append(("M54.1", "G54", WARN, "CR_3_5",
                   "Одночасне застосування M54.1 з кодами типу G54* не рекомендується (подвійне кодування)."))
for _a in ["M43.3", "M47.1", "M50.0"]:
    for _b in ["G99.2", "G95.9"]:
        EXCLUSIONS.append((_a, _b, WARN, "CR_4_3",
                           "Комбінований код уже містить мієлопатію — додаткове кодування мієлопатії не здійснюється."))
for _a in ["M32.1", "M34", "M35.3", "M10", "M11", "M00", "M01", "M02", "M07", "M45.0", "M46.1", "M05", "M06"]:
    EXCLUSIONS.append((_a, "M13", FAIL, "CR_6_3, табл. 1",
                       "Комбінований код уже включає артрит/артропатію — додаткове кодування артриту (M13) не здійснюється."))

# ----------------------------------------------------------------------------------------
# Розбір текстів
# ----------------------------------------------------------------------------------------
_CAT_RE = re.compile(r"Діагноз\s*_\s*(ПП|ФО|СФЗ|СП|С)(?![А-ЯІЇЄA-Z])")
_ST_RE = re.compile(r"([ST]\d{2}(?:\.\d)?)(?:\s*-\s*([ST]\d{2}(?:\.\d)?))?\.?-?")
_LEAD_RE = re.compile(r"^\s*([A-ZН]\d{2}(?:\.\d+)?)\.?\s*([*†])?")
_REF_RE = re.compile(r"([A-Z]\d{2}(?:\.\d+)?)(?:\s*-\s*([A-Z]\d{2}(?:\.\d+)?))?\s*\.?\s*-?\s*([*†])?")


def _clean(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").replace("\xa0", " ")).strip()


def _expand_st(a: str, b: str) -> list[str]:
    if not b:
        return [a]
    if "." not in a and "." not in b:
        return [f"{a[0]}{i:02d}" for i in range(int(a[1:3]), int(b[1:3]) + 1)]
    return [f"{a[:3]}.{i}" for i in range(int(a[4]), int(b[4]) + 1)]


def _expand_any(a: str, b: str) -> list[str]:
    if not b:
        return [a]
    if "." in a or "." in b:
        if a[:3] == b[:3]:
            return [f"{a[:3]}.{i}" for i in range(int(a[4]), int(b[4]) + 1)]
        a, b = a[:3], b[:3]
    out, letter = [], ord(a[0])
    while letter <= ord(b[0]):
        lo = int(a[1:3]) if letter == ord(a[0]) else 0
        hi = int(b[1:3]) if letter == ord(b[0]) else 99
        out += [f"{chr(letter)}{i:02d}" for i in range(lo, hi + 1)]
        letter += 1
    return out


def _readable(tok: str) -> str:
    return f"Діагноз_{tok}" if tok in CATS else tok


class TokenSet:
    """Набір токенів (категорії I–M і коди) з префіксами кодів для швидкого зіставлення."""
    __slots__ = ("exact", "prefixes")

    def __init__(self, tokens):
        self.exact = set(tokens)
        self.prefixes = set()
        for t in self.exact:
            if len(t) > 3 and t[1:3].isdigit():
                self.prefixes.add(t[:3])
                if len(t) > 5:
                    self.prefixes.add(t[:5])


def as_tokenset(tokens) -> TokenSet:
    return tokens if isinstance(tokens, TokenSet) else TokenSet(tokens)


def code_matches(code: str, ts: TokenSet) -> bool:
    """Код збігається з токеном, якщо один є уточненням іншого (S06 ↔ S06.20, R47 ↔ R47.0)."""
    return (code in ts.exact or code[:3] in ts.exact or code[:5] in ts.exact or code[:6] in ts.exact
            or code in ts.prefixes)


@dataclass
class Group:
    kind: str     # "cat" — категорії I–M; "code" — коди S/T (перемикач); "req" — інші коди (завжди обов'язкові);
                  # "case" — вимога до всього випадку
    tokens: list[str]

    @cached_property
    def ts(self) -> TokenSet:
        return TokenSet(self.tokens)

    @property
    def text(self) -> str:
        cats = [_readable(x) for x in self.tokens if x in CATS]
        codes = [x for x in self.tokens if x not in CATS]
        parts = cats + ([nszu.ranges_text(codes)] if len(codes) > 4 else codes)
        t = " або ".join(parts)
        if self.kind == "case":
            return t + " (обов'язковий у кожному випадку)"
        return t + (" (код первинної травми/опіку)" if self.kind == "code" else "")


# Обов'язкова для кожного випадку група: Діагноз_Z (коди Z50.x). СФЗ — окрема, необов'язкова категорія.
CASE_GROUPS: list[Group] = [Group("case", ["Z"])]
Z_PREFIX = "Z50"      # у довіднику без файлу категорій Z50.x позначено як СФЗ — вважаємо їх Діагноз_Z


def parse_formula(formula: str) -> tuple[list[Group], bool]:
    """
    Повертає групи формули та ознаку «*-код».
    «плюс» і «та код» розділяють обов'язкові групи; «та/або», «або» — варіанти в межах групи;
    у переліку кодів (без категорій) «та» також розділяє обов'язкові групи.
    """
    f = re.sub(r"^\s*плюс\s*", "", _clean(formula))
    parts: list[str] = []
    for p in re.split(r"\s*плюс\s*", f):
        parts += re.split(r"\s+та\s+(?=код)", p)
    groups, star = [], False
    for p in parts:
        if "*-код" in p:
            star, p = True, p.replace("*-код", "")
        cats = _CAT_RE.findall(p)
        if cats:
            codes: list[str] = []
            for a, b in _ST_RE.findall(_CAT_RE.sub("", p)):
                codes += _expand_st(a, b)
            groups.append(Group("cat", list(dict.fromkeys(cats + codes))))
            continue
        for q in re.split(r"\s+та\s+", p):          # перелік кодів: «S43.-, S53.- та T11.2» = (S43 | S53) і T11.2
            codes = []
            for a, b in _ST_RE.findall(q):
                codes += _expand_st(a, b)
            if codes:
                groups.append(Group("code", list(dict.fromkeys(codes))))
    return groups, star


# ----------------------------------------------------------------------------------------
# Модель
# ----------------------------------------------------------------------------------------
@dataclass
class Diagnosis:
    row: int
    name: str
    code: str
    mark: str                       # "*", "†" або ""
    cats: set[str]
    main_text: str                  # стовпець N
    formula: str                    # стовпець O
    groups: list[Group] = field(default_factory=list)
    pair_tokens: list[str] = field(default_factory=list)
    pair_text: str = ""
    pair_flag: int = 0              # 0 — не потрібно; 1 — обов'язково; 2 — ручна перевірка; 3 — рекомендовано
    # --- джерела НСЗУ (None — використовується довідник) ---
    main_conds: list[str] | None = None        # умови з «Переліку основних»; [] — не може бути основним
    comp_groups: list[Group] | None = None     # формула для ролі супутнього
    comp_formula: str = ""
    source: str = "довідник"

    @property
    def cgroups(self) -> list[Group]:
        """Формула для ролі супутнього (якщо окремої немає — та сама, що для основного)."""
        return self.comp_groups if self.comp_groups is not None else self.groups

    def _conds(self) -> list[str]:
        if self.main_conds is not None:
            return [c.lower() for c in self.main_conds]
        n = self.main_text.lower()
        return [n] if n.startswith("так") or n.startswith("допустимі") else []

    @staticmethod
    def _cond_in_period(n: str, period: str | None) -> bool:
        if period == PERIOD_LONG and "післягостр" in n:
            return False
        if period == PERIOD_POST and "довготрив" in n:
            return False
        return True

    @staticmethod
    def _cond_conditional(n: str) -> bool:
        return any(k in n for k in ("якщо", "допустим", "після операц"))

    @property
    def can_be_main(self) -> bool:
        return bool(self._conds())

    @property
    def conditional(self) -> bool:
        conds = self._conds()
        return bool(conds) and all(self._cond_conditional(n) for n in conds)

    def main_in_period(self, period: str | None) -> bool:
        return any(self._cond_in_period(n, period) for n in self._conds())

    @property
    def cats_text(self) -> str:
        return ", ".join(c for c in CATS if c in self.cats) or "немає в I–M"

    def matches(self, tokens) -> bool:
        """Чи підходить діагноз під набір токенів (категорії I–M або коди, у т. ч. уточнені/рубрики)."""
        ts = as_tokenset(tokens)
        if self.cats & ts.exact:
            return True
        return code_matches(self.code, ts)

    def all_tokens(self) -> list[str]:
        return [t for g in self.groups for t in g.tokens]

    def all_ctokens(self) -> list[str]:
        return [t for g in self.cgroups for t in g.tokens]


class Directory:
    """Довідник діагнозів із таблиці xlsx."""

    def __init__(self, path: str, sheet: str | None = None, nszu_files: "nszu.NszuFiles | str | None" = "auto"):
        """
        nszu_files: "auto" — шукати файли НСЗУ в папці довідника; None — лише довідник;
        або явний nszu.NszuFiles.
        """
        self.path = path
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else self._find_sheet(wb)
        rows = list(ws.iter_rows(values_only=True))
        hdr_i, col = self._find_header(rows)
        self.items: list[Diagnosis] = []
        seen: set[str] = set()
        raw = []
        for i in range(hdr_i + 1, len(rows)):
            r = rows[i]
            name = _clean(r[col["B"]] if col["B"] < len(r) else "")
            if not name or name in seen:
                continue
            seen.add(name)
            raw.append((i + 1, r, name))
        # коди з позначкою * (для визначення парних кодів)
        star_codes = set()
        for _, _, name in raw:
            m = _LEAD_RE.match(name)
            if m and m.group(2) == "*":
                star_codes.add(m.group(1))
        for rownum, r, name in raw:
            get = lambda k: _clean(r[col[k]]) if col[k] < len(r) and r[col[k]] is not None else ""
            m = _LEAD_RE.match(name)
            mark = (m.group(2) or "") if m else ""
            code = m.group(1) if m else name.split(" ")[0]
            cats = {c for c, k in zip(CATS, "IJKLM") if get(k).lower() == "так"}
            groups, star_formula = parse_formula(get("O"))
            d = Diagnosis(rownum, name, code, mark, cats, get("N"), get("O"), groups)
            self._set_pair(d, name[m.end():] if m else name, star_formula, star_codes)
            d.code = nszu.norm_code(d.code)
            if d.code.startswith(Z_PREFIX):
                d.cats = (d.cats - {"СФЗ"}) | {"Z"}
            self.items.append(d)
        self.ambulatory = load_ambulatory()
        if nszu_files == "auto":
            nszu_files = nszu.find_files(Path(path).parent)
        self.nszu = nszu_files if (nszu_files and nszu_files.any()) else None
        if self.nszu:
            self._apply_nszu(self.nszu)
        self.by_name = {d.name: d for d in self.items}
        self.by_code: dict[str, Diagnosis] = {}
        for d in self.items:
            self.by_code.setdefault(d.code, d)

    @property
    def sources(self) -> list[str]:
        return [Path(self.path).name] + (self.nszu.describe() if self.nszu else [])

    # --------------------------- файли НСЗУ ---------------------------
    @staticmethod
    def _groups(parsed) -> list[Group]:
        """Групи з булевої формули без «Діагноз_Z» (він — обов'язкова група випадку CASE_GROUPS)."""
        return [Group(kind, toks) for kind, toks in parsed if toks and toks != ["Z"]]

    def _apply_nszu(self, f: "nszu.NszuFiles"):
        main_list = nszu.load_main_list(f.main_list) if f.main_list else None
        fm = nszu.load_formulas(f.formulas_main) if f.formulas_main else {}
        fa = nszu.load_formulas(f.formulas_add) if f.formulas_add else {}
        categories = nszu.load_categories(f.categories) if f.categories else None
        universe = set(main_list or {}) | set(fm) | set(fa) | set(categories or {})
        # коди, що мають уточнення у файлах НСЗУ (рубрики G54, M54.1 …), окремими пунктами не показуються
        uprefix = {u[:n] for u in universe for n in (3, 5) if len(u) > n}
        leaves = universe - uprefix

        def lookup(table, code):
            """Формула коду; якщо коду немає — формула найближчої рубрики (M54.16 → M54.1 → M54)."""
            for k in (code, code[:5], code[:3]):
                if k in table:
                    return table[k]
            return None

        dir_codes: dict[str, Diagnosis] = {}
        for d in self.items:
            dir_codes.setdefault(d.code, d)
        items: list[Diagnosis] = []
        for d in self.items:
            if d.code in leaves:
                d.source = "НСЗУ"
                items.append(d)
            elif d.code in uprefix:
                continue                       # рубрика, замінена уточненими кодами з файлів НСЗУ
            else:
                items.append(d)                # є лише в довіднику
        def new_item(u: str) -> Diagnosis:
            parent = (dir_codes.get(u[:5]) if len(u) > 5 else None) or dir_codes.get(u[:3])
            if parent:
                title = re.sub(r"^\s*\S+\s*[*†]?\s*", "", parent.name).strip(" *†")
                return Diagnosis(0, f"{u} {title} (уточнення коду {parent.code})", u, "", set(parent.cats),
                                 "", "", [], source="НСЗУ")
            return Diagnosis(0, f"{u} (назви немає в довіднику)", u, "", set(), "", "", [], source="НСЗУ")

        for u in sorted(leaves - set(dir_codes)):
            items.append(new_item(u))
        # коди, названі у формулах (T11.2, S43.0x, †-коди G99.2 …), яких немає серед пунктів — щоб їх можна було обрати
        have = TokenSet([d.code for d in items])
        extra = set()
        for table in (fm, fa):
            for fr in table.values():
                for kind, toks in fr.groups:
                    if kind in ("code", "req"):
                        extra |= {t for t in toks if not (t in have.exact or t in have.prefixes
                                                         or t[:5] in have.exact or t[:3] in have.exact)}
        for u in sorted(extra):
            items.append(new_item(u))

        # Довідник дає лише назви й категорії I–M. Можливість бути основним, формули й пари —
        # тільки з файлів НСЗУ (для аспекту, файл якого є); стовпці N, O і позначки †/* у назвах не використовуються.
        for d in items:
            if categories is not None:          # категорії — лише з файлу категорій НСЗУ
                d.cats = set(categories.get(d.code, set()))
            if main_list is not None:
                d.main_conds = list(main_list.get(d.code, []))
                d.main_text = "; ".join(d.main_conds) if d.main_conds else "ні (немає в переліку основних НСЗУ)"
            if f.formulas_main or f.formulas_add:
                rm, ra = lookup(fm, d.code), lookup(fa, d.code)
                d.groups, d.formula = (self._groups(rm.groups), rm.text) if rm else ([], "")
                if ra:
                    d.comp_groups, d.comp_formula = self._groups(ra.groups), ra.text
                else:
                    d.comp_groups, d.comp_formula = d.groups, d.formula
        if f.pairs:
            self._apply_pairs(items, nszu.load_pairs(f.pairs))
        self.items = items

    @staticmethod
    def _apply_pairs(items: list[Diagnosis], pairs: list["nszu.PairRow"]):
        star_ts = [(p, TokenSet(p.stars)) for p in pairs]
        dag_ts = [(p, TokenSet(p.daggers)) for p in pairs]
        for d in items:
            srows = [p for p, ts in star_ts if code_matches(d.code, ts)]
            if srows:
                toks = [c for p in srows for c in p.daggers]
                d.pair_tokens, d.pair_flag = list(dict.fromkeys(toks)), 1
                d.pair_text = nszu.ranges_text(toks) + " †"
                continue
            drows = [p for p, ts in dag_ts if code_matches(d.code, ts)]
            if drows:                          # будь-який †-код з переліку пар обов'язково потребує *-коду
                toks = [c for p in drows for c in p.stars]
                d.pair_tokens, d.pair_flag = list(dict.fromkeys(toks)), 1
                d.pair_text = nszu.ranges_text(toks) + " *"
            else:                              # пари — лише з переліку пар (позначки †/* у назвах не враховуються)
                d.pair_tokens, d.pair_text, d.pair_flag = [], "", 0

    @staticmethod
    def _find_sheet(wb):
        for ws in wb.worksheets:
            for r in ws.iter_rows(min_row=1, max_row=10, values_only=True):
                if any(isinstance(v, str) and "Діагноз_ПП" in v for v in r):
                    return ws
        return wb.worksheets[0]

    @staticmethod
    def _find_header(rows):
        names = {"B": "Діагноз", "I": "Діагноз_ПП", "J": "Діагноз_ФО", "K": "Діагноз_С", "L": "Діагноз_СП",
                 "M": "Діагноз_СФЗ", "N": "Основний діагноз", "O": "Формула для кодування"}
        default = {k: ord(k) - ord("A") for k in names}
        for i, r in enumerate(rows[:20]):
            vals = [_clean(v) for v in r]
            if "Діагноз_ПП" in vals:
                col = dict(default)
                for k, h in names.items():
                    if h in vals:
                        col[k] = vals.index(h)
                return i, col
        return 3, default

    @staticmethod
    def _set_pair(d: Diagnosis, rest: str, star_formula: bool, star_codes: set[str]):
        """Парні коди †/* (розділ III Правил) з назви діагнозу."""
        startype = d.mark == "*" or (star_formula and d.mark != "†")
        toks, raw = [], []
        for a, b, mk in _REF_RE.findall(rest):
            if mk == "†" and not startype:
                continue
            if startype or mk == "*" or (not mk and a in star_codes):
                toks += _expand_any(a, b)
                raw.append(a + ("–" + b if b else "") + (mk or ("*" if a in star_codes else "")))
        d.pair_tokens = list(dict.fromkeys(toks))
        d.pair_text = ", ".join(raw) if raw else ""
        if d.mark in ("*", "†") or star_formula:
            d.pair_flag = 1 if toks else 2
        else:
            d.pair_flag = 3 if toks else 0

    # --------------------------- пошук ---------------------------
    @staticmethod
    def adhoc(text: str) -> Diagnosis | None:
        """Діагноз поза довідником, введений з кодом МКХ: враховується лише за кодом (пари †/*, переліки кодів)."""
        m = re.match(r"^\s*([A-ZА-ЯІ]\d{2}(?:\.\d{1,2})?)(?![\d.])", _clean(text))
        if not m:
            return None
        return Diagnosis(-1, _clean(text), nszu.norm_code(m.group(1)), "", set(), "", "", [], source="введено")

    def find(self, text: str) -> Diagnosis | None:
        """Точна назва або код (напр. «I63.3»)."""
        text = _clean(text)
        if text in self.by_name:
            return self.by_name[text]
        return self.by_code.get(nszu.norm_code(text))

    # --------------------------- списки ---------------------------
    def main_candidates(self, period: str | None) -> list[Diagnosis]:
        return [d for d in self.items if d.main_in_period(period)]

    @staticmethod
    def _formula_groups(main: Diagnosis, slot: int, selected: list[Diagnosis | None]):
        """Групи формул, що стосуються поля slot: групи основного, а для полів 2+ — ще й супутнього 1.
        Повертає [(власник, група)] без обов'язкових для всього випадку груп (CASE_GROUPS)."""
        groups = [(main, g) for g in main.groups]
        if slot > 0 and selected and selected[0] is not None:
            groups += [(selected[0], g) for g in selected[0].cgroups]
        return groups

    def companion_tokens(self, main: Diagnosis, slot: int, selected: list[Diagnosis | None],
                         narrow: bool = True) -> list[str]:
        """
        Токени для поля slot.
          * Групи, які власник формули закриває сам (своєю категорією), не враховуються;
            якщо сам діагноз закриває всі свої групи — вони лишаються як необов'язкові варіанти.
          * Обов'язкова група випадку (Діагноз_Z) додається завжди.
          * narrow=True (для випадаючого списку): лише групи, ще не закриті іншими обраними
            діагнозами; якщо все закрито — усі групи з попереднього пункту.
          * narrow=False (для перевірки): усі групи з першого пункту.
        Парні коди †/* основного та попередніх супутніх додаються завжди.
        """
        formula = self._formula_groups(main, slot, selected)
        own = [(o, g) for o, g in formula if not o.matches(g.ts)] or formula
        own = own + [(main, g) for g in CASE_GROUPS]
        use = own
        if narrow:
            others = [main] + [d for j, d in enumerate(selected) if j != slot and d is not None]
            # умови супутнього 1 закриваються лише супутніми (основний не враховується)
            unmet = [(o, g) for o, g in own
                     if not any(x is not o and not (o is not main and x is main) and x.matches(g.ts) for x in others)]
            if unmet:
                use = unmet
        tokens = [t for _, g in use for t in g.tokens] + main.pair_tokens
        for d in selected[:slot]:
            if d is not None:
                tokens += d.pair_tokens
        return tokens

    def companion_candidates(self, main: Diagnosis | None, slot: int,
                             selected: list[Diagnosis | None], narrow: bool = True) -> list[Diagnosis]:
        """slot — 0-базовий номер поля; selected — обрані супутні (у порядку полів)."""
        if main is None:
            return []
        ts = TokenSet(self.companion_tokens(main, slot, selected, narrow))
        excl = {main.name} | {d.name for d in selected[:slot] if d is not None}
        return [d for d in self.items if d.name not in excl and d.matches(ts)]


# ----------------------------------------------------------------------------------------
# План полів супутніх діагнозів (за ролями)
# ----------------------------------------------------------------------------------------
@dataclass
class Slot:
    """Поле супутнього діагнозу.
    role: "main" — вимога формули основного; "pair" — парний код †/*; "comp1" — вимога формули
    супутнього 1; "free" — необов'язкове (коморбідний стан). tokens=None — будь-який діагноз."""
    key: str
    role: str
    label: str
    tokens: list[str] | None
    value: str = ""
    diag: Diagnosis | None = None


ROLE_TITLES = {"main": "За формулою основного", "pair": "Пара †/*",
               "comp1": "За формулою супутнього 1", "free": "Додатковий (необов'язково)"}
MAX_SLOTS = 12


def plan_slots(directory: Directory, main: Diagnosis | None, values: dict[str, str]) -> list[Slot]:
    """
    Порядок полів:
      1) за формулою основного — по полю на кожну вимогу, яку основний не закриває сам (зокрема Діагноз_Z);
      2) пари †/* — для основного й діагнозів із полів 1), якщо пару ще не вказано;
      3) за формулою супутнього 1 (перше поле) — по полю на кожну вимогу, яку не закривають
         він сам і інші вже обрані супутні (основний не враховується);
      4) пари †/* для діагнозів із полів 3);
      5) необов'язкові поля (завжди одне порожнє в кінці).
    values — значення полів за ключами (ключі стабільні: m0, m1, p:main, p:m0, c0, p:c0, f0 …).
    """
    if main is None:
        return []

    def mk(key, role, label, tokens):
        v = values.get(key, "") or ""
        d = (directory.find(v) or directory.adhoc(v)) if _clean(v) else None
        return Slot(key, role, label, tokens, v, d)

    slots: list[Slot] = []
    groups = [g for g in main.groups if not main.matches(g.ts)]
    groups += [g for g in CASE_GROUPS if not main.matches(g.ts)]
    for i, g in enumerate(groups):
        slots.append(mk(f"m{i}", "main", f"{ROLE_TITLES['main']}: {g.text}", list(g.tokens)))
    main_slots = list(slots)
    free_vals = [mk(f"f{i}", "free", "", None) for i in range(MAX_SLOTS)]
    free_diags = [s.diag for s in free_vals if s.diag]

    def add_pairs(owners: list[tuple[str, Diagnosis]], pool: list[Diagnosis]):
        for key, owner in owners:
            if owner is None or owner.pair_flag != 1 or not owner.pair_tokens:
                continue
            ts = TokenSet(owner.pair_tokens)
            same = next((s for s in slots if s.role == "pair" and set(s.tokens) == set(owner.pair_tokens)), None)
            if same is not None:                     # та сама пара вже має поле — лише доповнюємо підпис
                if owner.code not in same.label:
                    same.label = same.label.replace(":", f", {owner.code}:", 1)
                continue
            needed = not any(x is not owner and x.matches(ts) for x in pool)
            if needed or _clean(values.get(f"p:{key}", "")):
                slots.append(mk(f"p:{key}", "pair", f"{ROLE_TITLES['pair']} для {owner.code}: {owner.pair_text}",
                                list(owner.pair_tokens)))

    base = [main] + [s.diag for s in main_slots if s.diag] + free_diags
    add_pairs([("main", main)] + [(s.key, s.diag) for s in main_slots], base)

    def will_cover(g: Group) -> bool:
        """Порожнє поле пари, усі варіанти якого закривають вимогу g (напр., G94.0 — ПП)."""
        for i, s in enumerate(slots):
            if s.role == "pair" and s.diag is None:
                cand = slot_candidates(directory, main, slots, i)
                if cand and all(d.matches(g.ts) for d in cand):
                    return True
        return False

    c1 = main_slots[0].diag if main_slots else None
    comp_slots: list[Slot] = []
    if c1 is not None:
        pool = [s.diag for s in slots if s.diag and s.key != "m0"] + free_diags     # без основного
        for k, g in enumerate(c1.cgroups):
            covered = c1.matches(g.ts) or any(x.matches(g.ts) for x in pool) or will_cover(g)
            if covered and not _clean(values.get(f"c{k}", "")):
                continue
            s = mk(f"c{k}", "comp1", f"{ROLE_TITLES['comp1']} ({c1.code}): {g.text}", list(g.tokens))
            slots.append(s)
            comp_slots.append(s)
        if comp_slots:
            base2 = [main] + [s.diag for s in slots if s.diag] + free_diags
            add_pairs([(s.key, s.diag) for s in comp_slots], base2)

    filled = [i for i, s in enumerate(free_vals) if s.diag or _clean(s.value)]
    n_free = min(MAX_SLOTS - len(slots), (max(filled) + 2) if filled else 1)
    for i in range(max(0, n_free)):
        s = free_vals[i]
        s.label = ROLE_TITLES["free"]
        slots.append(s)
    return slots


def slot_candidates(directory: Directory, main: Diagnosis | None, slots: list[Slot], index: int) -> list[Diagnosis]:
    """Варіанти для поля: діагнози, що відповідають ролі поля, без основного й обраних в інших полях."""
    if main is None:
        return []
    s = slots[index]
    taken = {main.name} | {x.diag.name for j, x in enumerate(slots) if j != index and x.diag}
    if s.tokens is None:
        return [d for d in directory.items if d.name not in taken]
    ts = TokenSet(s.tokens)
    return [d for d in directory.items if d.name not in taken and d.matches(ts)]


def slot_fits(slot: Slot) -> bool | None:
    """Чи відповідає обраний діагноз ролі поля (None — поле порожнє або без обмежень)."""
    if slot.diag is None or slot.tokens is None:
        return None
    return slot.diag.matches(TokenSet(slot.tokens))


# ----------------------------------------------------------------------------------------
# Перевірка
# ----------------------------------------------------------------------------------------
@dataclass
class Line:
    level: str      # ✔ ⚠ ✖ ℹ
    text: str


@dataclass
class Report:
    main: list[Line] = field(default_factory=list)
    main_groups: list[Line] = field(default_factory=list)
    companions: list[tuple[str, list[Line]]] = field(default_factory=list)
    exclusions: list[Line] = field(default_factory=list)
    summary: Line = field(default_factory=lambda: Line(INFO, ""))
    services: list[Line] = field(default_factory=list)      # лише для амбулаторії

    def all_lines(self) -> list[Line]:
        out = self.main + self.main_groups + self.exclusions
        for _, ls in self.companions:
            out += ls
        return out

    def as_text(self) -> str:
        s = ["ОСНОВНИЙ ДІАГНОЗ"] + [f"  {l.level} {l.text}" for l in self.main]
        s += ["", "ФОРМУЛА ОСНОВНОГО"] + [f"  {l.level} {l.text}" for l in self.main_groups]
        s += ["", "СУПУТНІ"]
        for i, (name, ls) in enumerate(self.companions, 1):
            s.append(f"  {i}. {name}")
            s += [f"     {l.level} {l.text}" for l in ls]
        s += ["", "НЕСУМІСНІ ПОЄДНАННЯ"] + ([f"  {l.level} {l.text}" for l in self.exclusions] or [f"  {OK} Не виявлено"])
        s += ["", f"ПІДСУМОК: {self.summary.level} {self.summary.text}"]
        if self.services:
            s += ["", "ПОСЛУГИ АМБУЛАТОРНОЇ РЕАБІЛІТАЦІЇ (за діагнозами)"] + [f"  {l.level} {l.text}" for l in self.services]
        return "\n".join(s)


def _group_met(g: Group, others: list[Diagnosis]) -> bool:
    return any(o.matches(g.ts) for o in others)


def _group_line(g: Group, met: bool, st_required: bool, prefix: str = "", self_d: Diagnosis | None = None) -> Line:
    if self_d is not None and self_d.matches(g.ts):
        own = ", ".join(c for c in CATS if c in self_d.cats and c in g.tokens) or self_d.code
        return Line(OK, f"{prefix}{g.text} — виконано самим діагнозом ({own}), додатковий не потрібен")
    if met:
        return Line(OK, f"{prefix}{g.text} — виконано")
    if g.kind == "code" and not st_required:
        return Line(WARN, f"{prefix}{g.text} — рекомендовано додати (CR_2_6, CR_5_6, CR_8_2, CR_11_2)")
    return Line(FAIL, f"{prefix}бракує: {g.text}")


def _pair_line(d: Diagnosis, others: list[Diagnosis]) -> Line | None:
    if d.pair_flag == 0:
        return None
    if d.pair_flag == 2:
        return Line(WARN, "Потрібен парний код †/* — у переліку пар не вказано який, перевірте вручну")
    if any(o.matches(d.pair_tokens) for o in others):
        return Line(OK, f"Парний код †/* є ({d.pair_text})")
    if d.pair_flag == 1:
        return Line(FAIL, f"Додайте парний код †/*: {d.pair_text}")
    return Line(WARN, f"Рекомендовано парний код: {d.pair_text}")


# ----------------------------------------------------------------------------------------
# Амбулаторна реабілітація: відповідність послугам АР1/АР2, АР3-1, АР3-2, АР4 за діагнозами
# (лист НСЗУ № 8575/8-15-26 від 01.04.2026; переліки — ambulatory_services.json)
# ----------------------------------------------------------------------------------------
def load_ambulatory(path: str | Path | None = None) -> dict | None:
    import json
    p = Path(path) if path else AMBULATORY_FILE
    if not p.is_file():
        return None
    data = json.loads(p.read_text(encoding="utf-8"))
    for s in data["services"]:
        s["codes"] = [nszu.norm_code(c) for c in s["codes"]]
    return data


def _coef(x: float) -> str:
    return f"{x:.1f}".replace(".", ",")


def service_label(s: dict, sr_record: bool) -> str:
    if s["id"] == "AP12":
        return f"АР1 (коеф. {_coef(s['coef_sr'])})" if sr_record else f"АР2 (коеф. {_coef(s['coef_no_sr'])})"
    return f"{s['name']} (коеф. {_coef(s['coef'])})"


def ambulatory_lines(amb: dict, main: Diagnosis, comps: list[Diagnosis], sr_record: bool) -> list[Line]:
    """
    Для кожної послуги: ✔ — основний діагноз у переліку; ⚠ — у переліку лише супутній (у листі —
    «щонайменше один з переліку», без уточнення, чи має це бути основний); ✖ — жодного.
    Інші критерії (направлення, епізод, ЕМЗ, кількість сесій) програма не перевіряє — лише нагадує.
    """
    out: list[Line] = []
    for s in amb["services"]:
        ts = TokenSet(s["codes"])
        label = service_label(s, sr_record)
        hits = [c.code for c in comps if code_matches(c.code, ts)]
        if code_matches(main.code, ts):
            out.append(Line(OK, f"{label}: основний діагноз {main.code} є в переліку"))
        elif hits:
            out.append(Line(WARN, f"{label}: у переліку лише супутній {', '.join(dict.fromkeys(hits))} — "
                                  f"уточніть, чи достатньо супутнього («щонайменше один з переліку»)"))
        else:
            out.append(Line(FAIL, f"{label}: жодного діагнозу випадку немає в переліку"))
            continue
        out += [Line(INFO, f"    {n}") for n in s.get("notes", [])]
    out += [Line(INFO, n) for n in amb.get("general_notes", [])]
    return out


def validate(directory: Directory, period: str, main_text: str, companion_texts: list[str],
             st_required: bool = False, setting: str = SETTING_INPATIENT, sr_record: bool = False) -> Report:
    """
    period          — PERIOD_POST або PERIOD_LONG;
    main_text       — назва або код основного діагнозу;
    companion_texts — назви/коди супутніх (порожні рядки ігноруються; діагноз поза довідником
                      вважається коморбідним станом і у формулах не враховується);
    st_required     — чи є коди S/T у формулах наслідків (T90–T95) обов'язковими;
    setting         — SETTING_INPATIENT або SETTING_OUTPATIENT (амбулаторія: + відповідність послугам АР);
    sr_record       — амбулаторія: є запис СР1/СР2/СР3 за останні 3 місяці (АР1 замість АР2).
    """
    rep = Report()
    main = directory.find(main_text) if main_text else None
    # позиції зберігаються: супутній 1 — це перше поле, навіть якщо воно порожнє
    comps_raw = list(companion_texts)
    comps = [(directory.find(c) or directory.adhoc(c)) if _clean(c) else None for c in comps_raw]

    if main is None:
        rep.main.append(Line(FAIL, "Основний діагноз не обрано або не знайдено в довіднику"))
        rep.summary = Line(FAIL, "Оберіть основний діагноз")
        return rep

    # --- основний
    rep.main.append(Line(INFO, f"{main.name}  |  категорії: {main.cats_text}"))
    rep.main.append(Line(INFO, f"Формула: {main.formula or '—'}"))
    if main.source == "НСЗУ" and main.row == 0:
        rep.main.append(Line(INFO, "Кода немає в довіднику — назву й категорії взято з батьківської рубрики"))
    if not main.can_be_main:
        rep.main.append(Line(FAIL, "Не може бути основним (стовпець N)"))
    elif not main.main_in_period(period):
        rep.main.append(Line(FAIL, f"Не може бути основним у періоді «{period}» (стовпець N: {main.main_text})"))
    elif main.conditional:
        rep.main.append(Line(WARN, f"Умовно допустимий: {main.main_text}"))
    else:
        rep.main.append(Line(OK, "Може бути основним"))

    known = [c for c in comps if c is not None]
    pl = _pair_line(main, known)
    if pl:
        rep.main.append(pl)

    # --- формула основного: групи закриваються супутніми
    for i, g in enumerate(main.groups + CASE_GROUPS, 1):
        rep.main_groups.append(_group_line(g, _group_met(g, known), st_required, f"Група {i}: ", main))

    # --- супутні
    seen: list[str] = []
    for i, (txt, d) in enumerate(zip(comps_raw, comps)):
        if not _clean(txt):
            continue
        lines: list[Line] = []
        if d is None:
            lines.append(Line(INFO, "Немає в довіднику (коморбідний стан) — у формулах не враховується"))
            rep.companions.append((_clean(txt), lines))
            continue
        if d.source == "введено":
            if d.matches(directory.companion_tokens(main, i, comps, narrow=False)):
                lines.append(Line(OK, f"Немає в довіднику — враховано за кодом {d.code}: відповідає формулі"))
            else:
                lines.append(Line(INFO, f"Немає в довіднику (коморбідний стан) — враховано лише за кодом {d.code}"))
            rep.companions.append((d.name, lines))
            continue
        lines.append(Line(INFO, f"Категорії: {d.cats_text}"))
        if d.name == main.name:
            lines.append(Line(FAIL, "Збігається з основним"))
        if d.name in seen:
            lines.append(Line(FAIL, "Повтор"))
        seen.append(d.name)

        allowed = directory.companion_candidates(main, i, comps, narrow=False)
        if any(a.name == d.name for a in allowed):
            src = "формулі основного" if i == 0 else "формулам основного та супутнього 1"
            lines.append(Line(OK, f"Відповідає {src}"))
        elif "СФЗ" in d.cats:
            lines.append(Line(INFO, "СФЗ — допустимий як додатковий, у формулах не враховується"))
        elif d.name != main.name:
            src = "формулою основного" if i == 0 else "формулами основного та супутнього 1"
            lines.append(Line(WARN, f"Не передбачено {src} (допустимо лише як супутній стан)"))

        others = [main] + [c for j, c in enumerate(comps) if j != i and c is not None]
        if i == 0:   # власна формула — лише для супутнього 1; основний діагноз її не закриває
            lines.append(Line(INFO, f"Формула: {(d.comp_formula or d.formula) or '—'}"))
            for k, g in enumerate(d.cgroups, 1):
                lines.append(_group_line(g, _group_met(g, others[1:]), st_required, f"Його група {k}: ", d))
        pl = _pair_line(d, others)
        if pl:
            lines.append(pl)
        rep.companions.append((d.name, lines))

    # --- несумісні поєднання
    codes = [main.code] + [c.code for c in known]
    for a, b, sev, rule, expl in EXCLUSIONS:
        if any(x.startswith(a) for x in codes) and any(x.startswith(b) for x in codes):
            rep.exclusions.append(Line(sev, f"{a} + {b} ({rule}): {expl}"))

    # --- підсумок
    levels = [l.level for l in rep.all_lines()]
    if setting == SETTING_OUTPATIENT and directory.ambulatory:
        rep.services = ambulatory_lines(directory.ambulatory, main, known, sr_record)

    if not main.main_in_period(period):
        rep.summary = Line(FAIL, "Основний діагноз некоректний")
    elif FAIL in levels:
        rep.summary = Line(FAIL, "Кодування не відповідає правилам — виправте пункти «✖»")
    elif WARN in levels:
        rep.summary = Line(WARN, "Кодування відповідає формулі, але є попередження — перевірте пункти «⚠»")
    else:
        rep.summary = Line(OK, "Кодування відповідає формулі та правилам")
    return rep


# ----------------------------------------------------------------------------------------
def _cli():
    ap = argparse.ArgumentParser(description="Перевірка кодування діагнозів реабілітаційного випадку")
    ap.add_argument("xlsx", help="файл довідника (Додаток до наказу)")
    ap.add_argument("--period", choices=["post", "long"], default="post",
                    help="post — післягострий, long — довготривалий")
    ap.add_argument("--main", required=True, help="основний діагноз (код або повна назва)")
    ap.add_argument("--comp", action="append", default=[], help="супутній діагноз (можна повторювати)")
    ap.add_argument("--st-required", action="store_true", help="коди S/T у формулах T9x — обов'язкові")
    ap.add_argument("--list", action="store_true", help="показати список для наступного супутнього поля")
    a = ap.parse_args()
    d = Directory(a.xlsx)
    period = PERIOD_POST if a.period == "post" else PERIOD_LONG
    print(validate(d, period, a.main, a.comp, a.st_required).as_text())
    if a.list:
        main = d.find(a.main)
        sel = [d.find(c) for c in a.comp]
        cand = d.companion_candidates(main, len(sel), sel)
        print(f"\nСписок для супутнього {len(sel) + 1}: {len(cand)} діагнозів")
        for x in cand[:30]:
            print("  ", x.name)
        if len(cand) > 30:
            print("   …")


if __name__ == "__main__":
    _cli()
